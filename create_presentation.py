from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create a presentation object
presentation = Presentation()

# Define a function to create a slide with title and content

def add_slide(title, content):
    slide_layout = presentation.slide_layouts[1]  # Title and Content layout
    slide = presentation.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    content_placeholder = slide.placeholders[1]
    
    title_placeholder.text = title
    content_placeholder.text = content
    
    # Set the background color
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)  # White background
    
    # Customize font color
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame'):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(18)  # Set font size
                    run.font.color.rgb = RGBColor(31, 78, 121)  # Blue color

# Title Slide
slide_layout = presentation.slide_layouts[0]  # Title Slide layout
slide = presentation.slides.add_slide(slide_layout)
slide.shapes.title.text = 'Solar Panel Dust Detection Project'
slide.placeholders[1].text = 'An Overview'

# Add slides
slides_content = [
    ('Introduction', 'An overview of solar panel efficiency and its challenges due to dust accumulation.'),
    ('Problem Statement', 'Dust accumulation reduces efficiency by X%. This project addresses this issue.'),
    ('Objectives', '- To detect dust on solar panels using imaging techniques.
- To propose a model for automatic detection.'),
    ('Literature Review', 'Review existing methods and technologies used for dust detection on solar panels.'),
    ('Dataset', 'Description of the dataset used for training and testing the model.'),
    ('Methodology Overview', 'Outline of the methods and techniques used in this project.'),
    ('Feature Extraction', 'Techniques used to extract features from images.'),
    ('SMOTE', 'Description of SMOTE for handling imbalanced datasets.'),
    ('Image Generation', 'Methods for generating synthetic images based on the dataset.'),
    ('Quality Validation', 'Methods to validate the quality of the generated images and model outputs.'),
    ('Model Architecture', 'Overview of the neural network architecture used in this project.'),
    ('Data Augmentation', 'Techniques used to augment the dataset for improved model training.'),
    ('Cross-Validation', 'Description of cross-validation applied to the model.'),
    ('Training Configuration', 'Details on model training configurations and parameters chosen.'),
    ('Evaluation Metrics', 'Metrics used to evaluate the model performance.'),
    ('Results Summary', 'Summary of results obtained from the model testing and validation.'),
    ('Advantages', 'Highlight the advantages of the proposed method over existing ones.'),
    ('Challenges', 'Discuss challenges faced during the project and proposed solutions.'),
    ('Future Work', 'Outline potential future directions and improvements.'),
    ('Conclusion', 'Final thoughts on the project and its contributions to the field.'),
    ('Thank You', 'Questions?')
]

for title, content in slides_content:
    add_slide(title, content)

# Save the presentation
presentation.save('solar_panel_dust_detection_presentation.pptx')